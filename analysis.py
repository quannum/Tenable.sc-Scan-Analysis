import logging
from collections import defaultdict
from dataclasses import dataclass

from constants import (
    DEFAULT_EXPECTED_SHEET,
    EXCLUDE,
    FALLBACK_EXPECTED_SHEET,
    INCLUDE,
    SHEET_EXPECTED_RANGE_COMPLIANCE,
    SHEET_EXPECTED_VS_ACTUAL,
    STATUS_GAP,
    STATUS_OK,
    STATUS_PARTIAL,
)
from scope_utils import (
    ParsedScope,
    merge_intervals,
    parse_scope_item,
    scope_contains,
    scope_intersects,
    scope_size,
    scope_to_interval,
    split_scope_items,
    subtract_intervals,
)

LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True)
class ScopeRecord:
    parsed: ParsedScope
    scan_name: str
    scope_item: str


@dataclass(frozen=True)
class ExcludedScopeRecord:
    parsed: ParsedScope
    scan_name: str
    asset_name: str
    scope_item: str


@dataclass(frozen=True)
class ExclusionImpact:
    asset: str
    scope: str
    loss: int


@dataclass(frozen=True)
class CoverageResult:
    environment: str | None
    location: str | None
    scope_item: str
    expected_size: int
    covered_count: int
    gap_count: int
    percent_lost: float
    total_included_ips: int
    exclusion_ip_total: int
    covered: str
    covering_scans: set[str]
    status: str
    reason: str
    required_scan: str
    required_scan_covered: str
    coverage_pct: float


def filter_scans(scans, config):
    if (
        not config.include_keywords
        and not config.exclude_keywords
        and config.filter_disabled_mode == "ALL"
    ):
        LOGGER.info("No filtering applied. Using all %s scans", len(scans))
        return scans

    filtered = []

    include_keywords = (
        config.include_keywords
        if config.case_sensitive
        else [keyword.lower() for keyword in config.include_keywords]
    )
    exclude_keywords = (
        config.exclude_keywords
        if config.case_sensitive
        else [keyword.lower() for keyword in config.exclude_keywords]
    )

    for scan in scans:
        name = scan.get("info", {}).get("name") or scan.get("name", "")
        compare_name = name if config.case_sensitive else name.lower()

        schedule = scan.get("schedule", {})
        enabled_raw = schedule.get("enabled")

        if isinstance(enabled_raw, str):
            enabled = enabled_raw.lower() == "true"
        elif isinstance(enabled_raw, bool):
            enabled = enabled_raw
        else:
            enabled = True

        if config.filter_disabled_mode == "ENABLED_ONLY" and not enabled:
            continue
        if config.filter_disabled_mode == "DISABLED_ONLY" and enabled:
            continue

        include_pass = True
        if include_keywords:
            if config.match_all_include:
                include_pass = all(
                    keyword in compare_name for keyword in include_keywords
                )
            else:
                include_pass = any(
                    keyword in compare_name for keyword in include_keywords
                )

        exclude_pass = True
        if exclude_keywords and any(
            keyword in compare_name for keyword in exclude_keywords
        ):
            exclude_pass = False

        if include_pass and exclude_pass:
            filtered.append(scan)

    LOGGER.info("Scans before filter: %s", len(scans))
    LOGGER.info("Scans after filter: %s", len(filtered))
    LOGGER.info("Disabled filter mode: %s", config.filter_disabled_mode)

    return filtered


def normalize_scope(
    normalized_ws, scan_name, asset_name, inclusion_type, defined_string
):
    for scope_item in split_scope_items(defined_string):
        normalized_ws.append([scan_name, asset_name, inclusion_type, scope_item])


def walk_combination(
    node, scan_name, scope_ws, normalized_ws, data_access, in_complement=False
):
    if not isinstance(node, dict):
        return

    operator = node.get("operator")
    if operator == "complement":
        in_complement = True

    if "id" in node and operator is None:
        asset_id = node.get("id")
        if asset_id in (-1, None):
            return

        asset = data_access.get_asset(asset_id)
        if not asset:
            LOGGER.warning(
                "Scan '%s' references missing asset id '%s' in combination asset",
                scan_name,
                asset_id,
            )
            return

        defined = asset.get("typeFields", {}).get("definedIPs")
        asset_name = asset.get("name") or f"Asset {asset_id}"
        inclusion_type = EXCLUDE if in_complement else INCLUDE

        if defined:
            scope_ws.append([scan_name, inclusion_type, "Asset", asset_name, defined])
            normalize_scope(
                normalized_ws, scan_name, asset_name, inclusion_type, defined
            )
        else:
            LOGGER.warning(
                "Combination asset '%s' on scan '%s' has no defined IPs",
                asset_name,
                scan_name,
            )

        return

    walk_combination(
        node.get("operand1"),
        scan_name,
        scope_ws,
        normalized_ws,
        data_access,
        in_complement,
    )
    walk_combination(
        node.get("operand2"),
        scan_name,
        scope_ws,
        normalized_ws,
        data_access,
        in_complement,
    )


def build_scope_sheets(scope_ws, normalized_ws, data_access, config):
    all_scans = data_access.get_scans()
    filtered_scans = filter_scans(all_scans, config)

    for scan in filtered_scans:
        scan_id = scan.get("id")
        scan_name = scan.get("name")
        if scan_id in (None, "") or not scan_name:
            LOGGER.warning("Skipping malformed scan record: %s", scan)
            continue

        details = data_access.get_scan_details(scan_id)
        if not details:
            LOGGER.warning("Skipping scan '%s' because details were not found", scan_name)
            continue

        ip_list = details.get("ipList")
        if ip_list and ip_list != "*":
            scope_ws.append([scan_name, INCLUDE, "Scan", "Direct IP List", ip_list])
            normalize_scope(normalized_ws, scan_name, "SCAN_IPLIST", INCLUDE, ip_list)

        for asset_ref in details.get("assets", []):
            asset_id = asset_ref.get("id")
            if asset_id in (None, ""):
                LOGGER.warning(
                    "Scan '%s' contains asset reference without id", scan_name
                )
                continue

            asset = data_access.get_asset(asset_id)
            if not asset:
                LOGGER.warning(
                    "Scan '%s' references missing asset id '%s'", scan_name, asset_id
                )
                continue

            asset_type = asset.get("type")

            if asset_type == "static":
                defined = asset.get("typeFields", {}).get("definedIPs")
                asset_name = asset.get("name") or f"Asset {asset_id}"

                if defined:
                    scope_ws.append([scan_name, INCLUDE, "Asset", asset_name, defined])
                    normalize_scope(
                        normalized_ws, scan_name, asset_name, INCLUDE, defined
                    )
                else:
                    LOGGER.warning(
                        "Static asset '%s' on scan '%s' has no defined IPs",
                        asset_name,
                        scan_name,
                    )

            elif asset_type == "combination":
                walk_combination(
                    asset.get("typeFields", {}).get("combinations", {}),
                    scan_name,
                    scope_ws,
                    normalized_ws,
                    data_access,
                )
            else:
                LOGGER.debug(
                    "Skipping unsupported asset type '%s' for asset '%s' on scan '%s'",
                    asset_type,
                    asset.get("name"),
                    scan_name,
                )


def build_coverage_data(normalized_ws):
    actual_scopes = []
    excluded_scopes = []
    actual_by_scan = defaultdict(list)
    excluded_by_scan = defaultdict(list)

    for row in normalized_ws.iter_rows(min_row=2, values_only=True):
        scan_name, asset_name, inclusion_type, scope_item = row
        if not scope_item:
            continue

        try:
            parsed = parse_scope_item(scope_item)
        except ValueError as exc:
            LOGGER.warning(
                "Skipping invalid normalized scope for scan '%s', asset '%s': %s",
                scan_name,
                asset_name,
                exc,
            )
            continue

        if inclusion_type == INCLUDE:
            entry = ScopeRecord(parsed=parsed, scan_name=scan_name, scope_item=scope_item)
            actual_scopes.append(entry)
            actual_by_scan[scan_name].append(entry)
        elif inclusion_type == EXCLUDE:
            entry = ExcludedScopeRecord(
                parsed=parsed,
                scan_name=scan_name,
                asset_name=asset_name,
                scope_item=scope_item,
            )
            excluded_scopes.append(entry)
            excluded_by_scan[scan_name].append(entry)

    return actual_scopes, excluded_scopes, actual_by_scan, excluded_by_scan


def resolve_expected_sheet(expected_workbook):
    if DEFAULT_EXPECTED_SHEET in expected_workbook.sheetnames:
        return expected_workbook[DEFAULT_EXPECTED_SHEET]
    if FALLBACK_EXPECTED_SHEET in expected_workbook.sheetnames:
        return expected_workbook[FALLBACK_EXPECTED_SHEET]

    available = ", ".join(expected_workbook.sheetnames)
    raise KeyError(
        "Expected scope workbook must contain either "
        f"'{DEFAULT_EXPECTED_SHEET}' or '{FALLBACK_EXPECTED_SHEET}'. Found: {available}"
    )


def validate_expected_row(row):
    if len(row) < 3:
        raise ValueError(
            "Expected range rows must include at least Scope Item, Location, and Environment columns"
        )

    scope_item = row[0]
    location = row[1] if len(row) > 1 else None
    environment = row[2] if len(row) > 2 else None
    required_scan = row[3] if len(row) > 3 else None

    if not scope_item:
        return None

    return scope_item, location, environment, required_scan


def build_expected_analysis_sheets(workbook):
    compare_ws = workbook.create_sheet(SHEET_EXPECTED_VS_ACTUAL)
    compare_ws.append(
        [
            "Environment",
            "Location",
            "Expected",
            "Expected IPs",
            "Net Covered IPs",
            "Net Portfolio Gap",
            "% Scan Coverage Suppressed",
            "Covered",
            "Scans Covering",
            "Status",
            "Reason",
            "Required Scan",
            "Required Scan Covered",
        ]
    )

    compliance_ws = workbook.create_sheet(SHEET_EXPECTED_RANGE_COMPLIANCE)
    compliance_ws.append(
        [
            "Environment",
            "Location",
            "Expected Range",
            "Expected IPs",
            "Covered IPs",
            "Gap IPs",
            "Coverage %",
        ]
    )

    return compare_ws, compliance_ws


def build_totals():
    return {
        "portfolio_expected_total": 0,
        "portfolio_covered_total": 0,
        "portfolio_gap_total": 0,
        "portfolio_exclusion_total": 0,
        "portfolio_included_total": 0,
    }


def determine_required_scan_coverage(required_scan, covering_scans):
    normalized_required_scan = str(required_scan or "").strip()
    if not normalized_required_scan:
        return "", ""

    if normalized_required_scan in covering_scans:
        return normalized_required_scan, "Yes"

    return normalized_required_scan, "No"


def build_exclusion_reason(exclusion_ip_total, relevant_exclusions):
    exclusion_lines = []
    for excluded_scan_name in sorted(relevant_exclusions):
        for entry in relevant_exclusions[excluded_scan_name]:
            exclusion_lines.append(
                f"{excluded_scan_name} | {entry.asset} | "
                f"{entry.scope} ({entry.loss} IPs)"
            )

    return f"Excluded {exclusion_ip_total} IPs:\n" + "\n".join(exclusion_lines)


def determine_coverage_status(covered_count, expected_size, exclusion_ip_total, relevant_exclusions):
    if covered_count == expected_size:
        return STATUS_OK, "Yes", "Fully contained by scan scope"

    if covered_count > 0:
        if exclusion_ip_total > 0:
            return (
                STATUS_PARTIAL,
                "Partial",
                build_exclusion_reason(exclusion_ip_total, relevant_exclusions),
            )

        return STATUS_PARTIAL, "Partial", "Partial coverage detected"

    return STATUS_GAP, "No", "No scan scope intersects expected range"


def collect_covering_scans(actual_scopes, expected):
    full_cover_scans = set()
    partial_scans = set()

    for actual_scope in actual_scopes:
        if scope_contains(actual_scope.parsed, expected):
            full_cover_scans.add(actual_scope.scan_name)
        elif scope_intersects(actual_scope.parsed, expected):
            partial_scans.add(actual_scope.scan_name)

    return full_cover_scans.union(partial_scans)


def calculate_scan_intervals(
    scan_name,
    expected,
    expected_start,
    expected_end,
    actual_by_scan,
    excluded_by_scan,
):
    included = []
    excluded = []
    relevant_exclusions = []
    exclusion_ip_total = 0

    for actual_scope in actual_by_scan[scan_name]:
        if not scope_intersects(actual_scope.parsed, expected):
            continue

        actual_start, actual_end = scope_to_interval(actual_scope.parsed)
        overlap_start = max(actual_start, expected_start)
        overlap_end = min(actual_end, expected_end)

        if overlap_start <= overlap_end:
            included.append((overlap_start, overlap_end))

    for excluded_scope in excluded_by_scan[scan_name]:
        if not scope_intersects(excluded_scope.parsed, expected):
            continue

        excluded_start, excluded_end = scope_to_interval(excluded_scope.parsed)
        overlap_start = max(excluded_start, expected_start)
        overlap_end = min(excluded_end, expected_end)

        if overlap_start > overlap_end:
            continue

        excluded.append((overlap_start, overlap_end))
        loss = overlap_end - overlap_start + 1
        relevant_exclusions.append(
            ExclusionImpact(
                asset=excluded_scope.asset_name,
                scope=excluded_scope.scope_item,
                loss=loss,
            )
        )
        exclusion_ip_total += loss

    included = merge_intervals(included)
    excluded = merge_intervals(excluded)

    included_ip_total = sum(end - start + 1 for start, end in included)
    net_intervals = subtract_intervals(included, excluded)

    return included_ip_total, net_intervals, relevant_exclusions, exclusion_ip_total


def calculate_coverage_result(
    scope_item,
    location,
    environment,
    required_scan,
    expected,
    actual_scopes,
    actual_by_scan,
    excluded_by_scan,
    exclusion_impact_by_scan,
):
    expected_size = scope_size(expected)
    expected_start, expected_end = scope_to_interval(expected)
    covering_scans = collect_covering_scans(actual_scopes, expected)

    relevant_exclusions = defaultdict(list)
    cover_intervals = []
    total_included_ips = 0
    exclusion_ip_total = 0

    for scan_name in covering_scans:
        included_ips, net_intervals, scan_exclusions, scan_excluded_ips = (
            calculate_scan_intervals(
                scan_name,
                expected,
                expected_start,
                expected_end,
                actual_by_scan,
                excluded_by_scan,
            )
        )

        total_included_ips += included_ips
        cover_intervals.extend(net_intervals)
        exclusion_ip_total += scan_excluded_ips
        if scan_excluded_ips:
            exclusion_impact_by_scan[scan_name] += scan_excluded_ips

        if scan_exclusions:
            relevant_exclusions[scan_name].extend(scan_exclusions)

    cover_intervals = merge_intervals(cover_intervals)
    covered_count = sum(end - start + 1 for start, end in cover_intervals)
    covered_count = min(covered_count, expected_size)
    gap_count = max(0, expected_size - covered_count)
    percent_lost = (
        round((exclusion_ip_total / total_included_ips) * 100, 2)
        if total_included_ips
        else 0.0
    )

    status, covered, reason = determine_coverage_status(
        covered_count, expected_size, exclusion_ip_total, relevant_exclusions
    )
    required_scan_value, required_scan_covered = determine_required_scan_coverage(
        required_scan, covering_scans
    )
    coverage_pct = round((covered_count / expected_size) * 100, 2) if expected_size else 0.0

    return CoverageResult(
        environment=environment,
        location=location,
        scope_item=scope_item,
        expected_size=expected_size,
        covered_count=covered_count,
        gap_count=gap_count,
        percent_lost=percent_lost,
        total_included_ips=total_included_ips,
        exclusion_ip_total=exclusion_ip_total,
        covered=covered,
        covering_scans=covering_scans,
        status=status,
        reason=reason,
        required_scan=required_scan_value,
        required_scan_covered=required_scan_covered,
        coverage_pct=coverage_pct,
    )


def update_totals(totals, result):
    totals["portfolio_expected_total"] += result.expected_size
    totals["portfolio_covered_total"] += result.covered_count
    totals["portfolio_gap_total"] += result.gap_count
    totals["portfolio_exclusion_total"] += result.exclusion_ip_total
    totals["portfolio_included_total"] += result.total_included_ips


def append_coverage_result(compare_ws, compliance_ws, result):
    compare_ws.append(
        [
            result.environment,
            result.location,
            result.scope_item,
            result.expected_size,
            result.covered_count,
            result.gap_count,
            result.percent_lost,
            result.covered,
            ", ".join(sorted(result.covering_scans)),
            result.status,
            result.reason,
            result.required_scan,
            result.required_scan_covered,
        ]
    )

    compliance_ws.append(
        [
            result.environment,
            result.location,
            result.scope_item,
            result.expected_size,
            result.covered_count,
            result.gap_count,
            result.coverage_pct,
        ]
    )


def analyze_expected_ranges(
    workbook, expected_scope_file, actual_scopes, actual_by_scan, excluded_by_scan
):
    compare_ws = None
    compliance_ws = None
    exclusion_impact_by_scan = defaultdict(int)
    totals = build_totals()

    if not expected_scope_file:
        LOGGER.info(
            "No expected scope workbook selected; skipping expected-vs-actual analysis"
        )
        return compare_ws, compliance_ws, exclusion_impact_by_scan, totals

    from openpyxl import load_workbook

    expected_wb = load_workbook(expected_scope_file)
    expected_ws = resolve_expected_sheet(expected_wb)
    compare_ws, compliance_ws = build_expected_analysis_sheets(workbook)

    for row_index, row in enumerate(
        expected_ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        try:
            validated_row = validate_expected_row(row)
        except ValueError as exc:
            LOGGER.warning("Skipping invalid expected range row %s: %s", row_index, exc)
            continue

        if validated_row is None:
            continue

        scope_item, location, environment, required_scan = validated_row

        try:
            expected = parse_scope_item(scope_item)
        except ValueError as exc:
            LOGGER.warning(
                "Skipping invalid expected scope in row %s: %s", row_index, exc
            )
            continue

        result = calculate_coverage_result(
            scope_item=scope_item,
            location=location,
            environment=environment,
            required_scan=required_scan,
            expected=expected,
            actual_scopes=actual_scopes,
            actual_by_scan=actual_by_scan,
            excluded_by_scan=excluded_by_scan,
            exclusion_impact_by_scan=exclusion_impact_by_scan,
        )
        update_totals(totals, result)
        append_coverage_result(compare_ws, compliance_ws, result)

    return compare_ws, compliance_ws, exclusion_impact_by_scan, totals
