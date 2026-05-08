from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .constants import (
    EXCLUDE,
    GRAY,
    GREEN,
    HEADER_FONT,
    INCLUDE,
    RED,
    SHEET_EXECUTIVE_SUMMARY,
    SHEET_RUN_METADATA,
    SHEET_SCAN_SCOPE_NORMALIZED,
    SHEET_SCAN_SCOPE_SUMMARY,
    SHEET_TOP_EXCLUSION_IMPACT,
    SHEET_WARNINGS,
    STATUS_GAP,
    STATUS_OK,
    STATUS_PARTIAL,
    VERSION,
    YELLOW,
)


def build_workbook():
    workbook = Workbook()
    scope_ws = workbook.create_sheet(SHEET_SCAN_SCOPE_SUMMARY)
    normalized_ws = workbook.create_sheet(SHEET_SCAN_SCOPE_NORMALIZED)
    workbook.remove(workbook["Sheet"])

    scope_ws.append(
        [
            "Scan Name",
            "Inclusion Type",
            "Source Type",
            "Source Name",
            "Scope Definition",
        ]
    )
    normalized_ws.append(["Scan Name", "Asset Name", "Inclusion Type", "Scope Item"])

    return workbook, scope_ws, normalized_ws


def append_executive_summary(workbook, totals):
    exec_ws = workbook.create_sheet(SHEET_EXECUTIVE_SUMMARY)
    exec_ws.append(["Metric", "Value", "Note"])

    overall_coverage_pct = (
        round(
            (totals.portfolio_covered_total / totals.portfolio_expected_total)
            * 100,
            2,
        )
        if totals.portfolio_expected_total
        else 0.0
    )

    portfolio_exclusion_loss_pct = (
        round(
            (totals.portfolio_exclusion_total / totals.portfolio_included_total)
            * 100,
            2,
        )
        if totals.portfolio_included_total
        else 0.0
    )

    exec_ws.append(
        [
            "Total Expected IPs",
            totals.portfolio_expected_total,
            "Number of IP addresses from expected ranges (Global IP Address Trackers)",
        ]
    )
    exec_ws.append(
        [
            "Total Net Covered IPs",
            totals.portfolio_covered_total,
            "Number of IP addresses covered by scans",
        ]
    )
    exec_ws.append(
        [
            "Total Gap IPs",
            totals.portfolio_gap_total,
            "Number of IP addresses not covered by scans",
        ]
    )
    exec_ws.append(
        [
            "Overall Portfolio Coverage %",
            overall_coverage_pct,
            "Percentage of IP addresses covered by scans (Total Net Covered IPs/Total Expected IPs)",
        ]
    )
    exec_ws.append(
        [
            "Total IPs Lost Due To Exclusions",
            totals.portfolio_exclusion_total,
            "Number of IP addresses excluded in scans",
        ]
    )
    exec_ws.append(
        [
            "% Coverage Lost Due To Exclusions",
            portfolio_exclusion_loss_pct,
            "Percentage of coverage lost due to exclusions (Total Excluded/Total Included)",
        ]
    )

    workbook.move_sheet(exec_ws, offset=1)
    return exec_ws


def build_impact_sheet(workbook, exclusion_impact_by_scan):
    impact_ws = workbook.create_sheet(SHEET_TOP_EXCLUSION_IMPACT)
    impact_ws.append(["Scan Name", "Total IPs Excluded"])

    for scan_name, total in sorted(
        exclusion_impact_by_scan.items(), key=lambda item: item[1], reverse=True
    ):
        impact_ws.append([scan_name, total])

    return impact_ws


def build_warning_sheet(workbook, warning_records):
    if not warning_records:
        return None

    warning_ws = workbook.create_sheet(SHEET_WARNINGS)
    warning_ws.append(["Level", "Logger", "Message"])

    for record in warning_records:
        warning_ws.append(
            [record.get("level"), record.get("logger"), record.get("message")]
        )

    return warning_ws


def build_run_metadata_sheet(workbook, config, output_file):
    metadata_ws = workbook.create_sheet(SHEET_RUN_METADATA)
    metadata_ws.append(["Field", "Value"])

    run_started_at = getattr(config, "run_started_at", None)
    run_started_value = (
        run_started_at.isoformat(sep=" ", timespec="seconds")
        if run_started_at
        else ""
    )

    rows = [
        ("Run Started At", run_started_value),
        ("Tool Version", VERSION),
        ("Mode", config.mode),
        ("Scan JSON Dir", config.scan_json_dir or ""),
        ("Asset JSON Dir", config.asset_json_dir or ""),
        ("Expected Scope File", config.expected_scope_file or ""),
        ("Expected Sheet", config.expected_sheet or ""),
        ("Output File", str(output_file)),
        ("Log File", str(config.log_file) if config.log_file else ""),
        ("Include Keywords", ", ".join(config.include_keywords)),
        ("Exclude Keywords", ", ".join(config.exclude_keywords)),
        ("Match All Include", str(config.match_all_include)),
        ("Case Sensitive", str(config.case_sensitive)),
        ("Filter Disabled Mode", config.filter_disabled_mode),
    ]

    for row in rows:
        metadata_ws.append(row)

    return metadata_ws


def find_column(ws, header_name):
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == header_name:
            return idx
    return None


def auto_wrap_and_adjust(ws, column_name):
    column_index = find_column(ws, column_name)
    if not column_index:
        return

    for row in ws.iter_rows(min_row=2):
        cell = row[column_index - 1]
        if not cell.value:
            continue

        cell.alignment = Alignment(wrap_text=True)
        line_count = str(cell.value).count("\n") + 1
        ws.row_dimensions[cell.row].height = 15 * line_count


def format_sheet(
    ws, status_col=None, include_exclude_col=None, percent_col=None, value_col=None
):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = GRAY

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 60)

    if status_col:
        for row in ws.iter_rows(min_row=2):
            cell = row[status_col - 1]
            if cell.value == STATUS_OK:
                cell.fill = GREEN
            elif cell.value == STATUS_PARTIAL:
                cell.fill = YELLOW
            elif cell.value == STATUS_GAP:
                cell.fill = RED

    if include_exclude_col:
        for row in ws.iter_rows(min_row=2):
            cell = row[include_exclude_col - 1]
            if cell.value == INCLUDE:
                cell.fill = GREEN
            elif cell.value == EXCLUDE:
                cell.fill = RED

    if percent_col:
        for row in ws.iter_rows(min_row=2):
            cell = row[percent_col - 1]
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
                if cell.value >= 90:
                    cell.fill = GREEN
                elif cell.value >= 70:
                    cell.fill = YELLOW
                else:
                    cell.fill = RED

    if value_col:
        for row in ws.iter_rows(min_row=2):
            metric_cell = row[0]
            target_cell = row[value_col - 1]

            if "Coverage %" in str(metric_cell.value):
                target_cell.number_format = "0.00"
                if isinstance(target_cell.value, (int, float)):
                    if target_cell.value >= 95:
                        target_cell.fill = GREEN
                    elif target_cell.value >= 85:
                        target_cell.fill = YELLOW
                    else:
                        target_cell.fill = RED

            if metric_cell.value == "Total Gap IPs" and isinstance(
                target_cell.value, (int, float)
            ):
                if target_cell.value > 0:
                    target_cell.fill = RED


def format_workbook(
    scope_ws,
    normalized_ws,
    compare_ws,
    compliance_ws,
    impact_ws,
    exec_ws,
    warning_ws=None,
    metadata_ws=None,
):
    format_sheet(scope_ws, include_exclude_col=find_column(scope_ws, "Inclusion Type"))

    if compare_ws:
        format_sheet(compare_ws, status_col=find_column(compare_ws, "Status"))
        auto_wrap_and_adjust(compare_ws, "Reason")

    if compliance_ws:
        format_sheet(
            compliance_ws, percent_col=find_column(compliance_ws, "Coverage %")
        )

    format_sheet(impact_ws)
    format_sheet(exec_ws, value_col=find_column(exec_ws, "Value"))
    if warning_ws:
        format_sheet(warning_ws)
    if metadata_ws:
        format_sheet(metadata_ws)

    normalized_ws.sheet_state = "hidden"
