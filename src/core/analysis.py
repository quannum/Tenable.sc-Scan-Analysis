import logging
from collections import defaultdict
from dataclasses import dataclass

from ..constants import (
    DEFAULT_EXPECTED_SHEET,
    FALLBACK_EXPECTED_SHEET,
    SHEET_EXPECTED_RANGE_COMPLIANCE,
    SHEET_EXPECTED_VS_ACTUAL,
)
from .scope_utils import parse_scope_item
from .tenable_scope_analysis import (
    CoverageResult,
    ExcludedScopeRecord,
    ExclusionImpact,
    ScopeRecord,
    build_coverage_data,
    build_scope_sheets,
    calculate_coverage_result,
    calculate_scan_intervals,
    collect_covering_scans,
    determine_coverage_status,
    determine_required_scan_coverage,
    filter_scans,
    normalize_scope,
    walk_combination,
)

LOGGER = logging.getLogger(__name__)

__all__ = [
    "CoverageResult",
    "ExcludedScopeRecord",
    "ExclusionImpact",
    "ScopeRecord",
    "Totals",
    "analyze_expected_ranges",
    "append_coverage_result",
    "build_coverage_data",
    "build_expected_analysis_sheets",
    "build_scope_sheets",
    "build_totals",
    "calculate_coverage_result",
    "calculate_scan_intervals",
    "collect_covering_scans",
    "determine_coverage_status",
    "determine_required_scan_coverage",
    "filter_scans",
    "normalize_scope",
    "resolve_expected_sheet",
    "update_totals",
    "validate_expected_row",
    "walk_combination",
]


@dataclass
class Totals:
    portfolio_expected_total: int = 0
    portfolio_covered_total: int = 0
    portfolio_gap_total: int = 0
    portfolio_exclusion_total: int = 0
    portfolio_included_total: int = 0


def resolve_expected_sheet(expected_workbook, expected_sheet=None):
    if expected_sheet:
        if expected_sheet in expected_workbook.sheetnames:
            return expected_workbook[expected_sheet]

        lowered = expected_sheet.lower()
        case_insensitive_matches = [
            sheet_name
            for sheet_name in expected_workbook.sheetnames
            if sheet_name.lower() == lowered
        ]
        if len(case_insensitive_matches) == 1:
            return expected_workbook[case_insensitive_matches[0]]
        if len(case_insensitive_matches) > 1:
            matches = ", ".join(case_insensitive_matches)
            raise KeyError(
                f"Requested expected sheet '{expected_sheet}' matched multiple sheets "
                f"case-insensitively: {matches}"
            )

        available = ", ".join(expected_workbook.sheetnames)
        raise KeyError(
            f"Expected scope workbook does not contain requested sheet "
            f"'{expected_sheet}'. Found: {available}"
        )

    if DEFAULT_EXPECTED_SHEET in expected_workbook.sheetnames:
        return expected_workbook[DEFAULT_EXPECTED_SHEET]
    if FALLBACK_EXPECTED_SHEET in expected_workbook.sheetnames:
        return expected_workbook[FALLBACK_EXPECTED_SHEET]

    available = ", ".join(expected_workbook.sheetnames)
    raise KeyError(
        "Expected scope workbook must contain either "
        f"'{DEFAULT_EXPECTED_SHEET}' or '{FALLBACK_EXPECTED_SHEET}'. Found: {available}"
    )


def validate_expected_row(row):
    if len(row) < 3:
        raise ValueError(
            "Expected range rows must include at least Scope Item, "
            "Location, and Environment columns"
        )

    scope_item = row[0]
    location = row[1] if len(row) > 1 else None
    environment = row[2] if len(row) > 2 else None
    required_scan = row[3] if len(row) > 3 else None

    if not scope_item:
        return None

    return scope_item, location, environment, required_scan


def build_expected_analysis_sheets(workbook):
    compare_ws = workbook.create_sheet(SHEET_EXPECTED_VS_ACTUAL)
    compare_ws.append(
        [
            "Environment",
            "Location",
            "Expected",
            "Expected IPs",
            "Net Covered IPs",
            "Net Portfolio Gap",
            "% Scan Coverage Suppressed",
            "Covered",
            "Scans Covering",
            "Status",
            "Reason",
            "Required Scan",
            "Required Scan Covered",
        ]
    )

    compliance_ws = workbook.create_sheet(SHEET_EXPECTED_RANGE_COMPLIANCE)
    compliance_ws.append(
        [
            "Environment",
            "Location",
            "Expected Range",
            "Expected IPs",
            "Covered IPs",
            "Gap IPs",
            "Coverage %",
        ]
    )

    return compare_ws, compliance_ws


def build_totals():
    return Totals()


def update_totals(totals, result):
    totals.portfolio_expected_total += result.expected_size
    totals.portfolio_covered_total += result.covered_count
    totals.portfolio_gap_total += result.gap_count
    totals.portfolio_exclusion_total += result.exclusion_ip_total
    totals.portfolio_included_total += result.total_included_ips


def append_coverage_result(compare_ws, compliance_ws, result):
    compare_ws.append(
        [
            result.environment,
            result.location,
            result.scope_item,
            result.expected_size,
            result.covered_count,
            result.gap_count,
            result.percent_lost,
            result.covered,
            ", ".join(sorted(result.covering_scans)),
            result.status,
            result.reason,
            result.required_scan,
            result.required_scan_covered,
        ]
    )

    compliance_ws.append(
        [
            result.environment,
            result.location,
            result.scope_item,
            result.expected_size,
            result.covered_count,
            result.gap_count,
            result.coverage_pct,
        ]
    )


def analyze_expected_ranges(
    workbook,
    expected_scope_file,
    actual_scopes,
    actual_by_scan,
    excluded_by_scan,
    expected_sheet=None,
):
    compare_ws = None
    compliance_ws = None
    exclusion_impact_by_scan = defaultdict(int)
    totals = build_totals()

    if not expected_scope_file:
        LOGGER.info(
            "No expected scope workbook selected; skipping expected-vs-actual analysis"
        )
        return compare_ws, compliance_ws, exclusion_impact_by_scan, totals

    from openpyxl import load_workbook

    expected_wb = load_workbook(expected_scope_file)
    expected_ws = resolve_expected_sheet(expected_wb, expected_sheet)
    compare_ws, compliance_ws = build_expected_analysis_sheets(workbook)

    for row_index, row in enumerate(
        expected_ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        try:
            validated_row = validate_expected_row(row)
        except ValueError as exc:
            LOGGER.warning("Skipping invalid expected range row %s: %s", row_index, exc)
            continue

        if validated_row is None:
            continue

        scope_item, location, environment, required_scan = validated_row

        try:
            expected = parse_scope_item(scope_item)
        except ValueError as exc:
            LOGGER.warning(
                "Skipping invalid expected scope in row %s: %s", row_index, exc
            )
            continue

        result = calculate_coverage_result(
            scope_item=scope_item,
            location=location,
            environment=environment,
            required_scan=required_scan,
            expected=expected,
            actual_scopes=actual_scopes,
            actual_by_scan=actual_by_scan,
            excluded_by_scan=excluded_by_scan,
            exclusion_impact_by_scan=exclusion_impact_by_scan,
        )
        update_totals(totals, result)
        append_coverage_result(compare_ws, compliance_ws, result)

    return compare_ws, compliance_ws, exclusion_impact_by_scan, totals
