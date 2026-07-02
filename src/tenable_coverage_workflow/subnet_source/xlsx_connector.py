import ipaddress
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..models import (
    CoverageTarget,
    NetworkRange,
    PrivateNetworkRange,
    SiteNetworkDefinition,
    SourceLoadResult,
    ValidationIssue,
    VlanRange,
)
from .validation import add_relationship_issues

HEADER_ALIASES = {
    "scope": {"scope item", "scope", "cidr", "ip range", "network"},
    "site_code": {"site code", "site_code", "code"},
    "site_name": {"site name", "site_name", "name", "location"},
    "location": {"location"},
    "description": {"description"},
    "region": {"region"},
    "timezone": {"timezone", "time zone"},
    "tags": {"tags"},
    "environment": {"environment"},
    "business_function": {"business function", "business_function", "function"},
    "target_type": {"target type", "target_type", "range type"},
    "vlan_name": {"vlan name", "vlan_name"},
    "vlan_tag": {"vlan id", "vlan tag", "vlan_id", "vlan_tag"},
    "required_scan": {"required scan", "scan name"},
    "required_asset": {"required asset", "asset name", "asset group"},
    "required_policy": {"required policy", "policy name"},
}


def load_xlsx_definitions(
    path_value: str | Path, sheet_name: str | None = None, audit_logger=None
) -> SourceLoadResult:
    path = Path(path_value)
    result = SourceLoadResult(files_processed=1)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        result.files_failed = 1
        result.validation_issues.append(ValidationIssue(str(path), str(exc)))
        return result
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            result.files_failed = 1
            result.validation_issues.append(
                ValidationIssue(
                    str(path),
                    f"Worksheet '{sheet_name}' was not found. Available: "
                    + ", ".join(workbook.sheetnames),
                )
            )
            workbook.close()
            return result
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook[workbook.sheetnames[0]]

    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        result.files_failed = 1
        result.validation_issues.append(
            ValidationIssue(str(path), "XLSX worksheet is empty.")
        )
        workbook.close()
        return result
    columns = _map_headers(headers)
    if "scope" not in columns:
        result.files_failed = 1
        result.validation_issues.append(
            ValidationIssue(
                str(path),
                "XLSX requires a Scope Item, Scope, CIDR, IP Range, or Network column.",
            )
        )
        workbook.close()
        return result

    site_targets: dict[str, list[CoverageTarget]] = defaultdict(list)
    for row_number, values in enumerate(rows, start=2):
        row = {
            name: values[index] if index < len(values) else None
            for name, index in columns.items()
        }
        scope = _text(row.get("scope"))
        if not scope:
            continue
        source_file = f"{path}#{worksheet.title}!{row_number}"
        try:
            networks = _parse_scope(scope)
        except ValueError as exc:
            result.validation_issues.append(
                ValidationIssue(source_file, f"Invalid scope '{scope}': {exc}")
            )
            continue
        site_name = _text(row.get("site_name")) or _text(row.get("location"))
        site_code = _text(row.get("site_code")) or _derive_site_code(site_name)
        if not site_code:
            result.validation_issues.append(
                ValidationIssue(source_file, "Site Code or Location is required.")
            )
            continue
        tags = _tags(row.get("tags"))
        for network in networks:
            target_type = _target_type(row, network)
            target = CoverageTarget(
                target_type=target_type,
                cidr=str(network),
                site_code=site_code,
                site_name=site_name or site_code,
                location=_optional(row.get("location")) or site_name or None,
                region=_optional(row.get("region")),
                description=_optional(row.get("description")),
                vlan_name=_optional(row.get("vlan_name")),
                vlan_tag=row.get("vlan_tag"),
                source_file=source_file,
                required_asset_name=_optional(row.get("required_asset")),
                required_scan_name=_optional(row.get("required_scan")),
                required_policy_name=_optional(row.get("required_policy")),
                timezone=_optional(row.get("timezone")),
                tags=tags,
                environment=_optional(row.get("environment")),
                business_function=_optional(row.get("business_function")),
            )
            result.coverage_targets.append(target)
            site_targets[site_code].append(target)
            if audit_logger:
                audit_logger.emit(
                    "coverage_target_created",
                    source_file=source_file,
                    site_code=site_code,
                    target_type=target_type,
                    cidr=str(network),
                )

    result.site_definitions.extend(
        _build_site_definition(site_code, targets)
        for site_code, targets in sorted(site_targets.items())
    )
    if not result.coverage_targets:
        result.files_failed = 1
    workbook.close()
    return add_relationship_issues(result)


def _build_site_definition(
    site_code: str, targets: list[CoverageTarget]
) -> SiteNetworkDefinition:
    first = targets[0]
    public_ranges = []
    private_ranges = []
    for target in targets:
        network = ipaddress.ip_network(target.cidr)
        base = NetworkRange(
            name=target.vlan_name,
            description=target.description,
            cidr=str(network),
            network=str(network.network_address),
            prefix_length=network.prefixlen,
            subnetmask=str(network.netmask),
        )
        if target.target_type == "PUBLIC":
            public_ranges.append(base)
            continue
        vlans = []
        if target.target_type == "VLAN":
            vlans.append(
                VlanRange(
                    name=target.vlan_name or "VLAN",
                    vlan_tag=target.vlan_tag,
                    description=target.description,
                    cidr=base.cidr,
                    network=base.network,
                    prefix_length=base.prefix_length,
                    subnetmask=base.subnetmask,
                )
            )
        private_ranges.append(PrivateNetworkRange(**base.__dict__, vlans=vlans))
    return SiteNetworkDefinition(
        source_file=str(first.source_file),
        site_name=first.site_name or site_code,
        site_code=site_code,
        description=first.description,
        location=first.location,
        region=first.region,
        timezone=first.timezone,
        tags=list(first.tags),
        environment=first.environment,
        business_function=first.business_function,
        scan_classification=dict(first.scan_classification),
        public_ranges=public_ranges,
        private_ranges=private_ranges,
    )


def _map_headers(headers) -> dict[str, int]:
    normalized = {
        _normalize_header(value): index for index, value in enumerate(headers)
    }
    result = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[field] = normalized[alias]
                break
    return result


def _parse_scope(value: str) -> list[ipaddress.IPv4Network]:
    if "-" in value:
        start_text, end_text = value.split("-", 1)
        start = ipaddress.ip_address(start_text.strip())
        end = ipaddress.ip_address(end_text.strip())
        if start.version != 4 or end.version != 4:
            raise ValueError("IPv6 is not supported")
        return list(ipaddress.summarize_address_range(start, end))
    network = ipaddress.ip_network(value, strict=False)
    if network.version != 4:
        raise ValueError("IPv6 is not supported")
    return [network]


def _target_type(row: dict[str, Any], network) -> str:
    explicit = _text(row.get("target_type")).upper().replace(" ", "_")
    aliases = {
        "PRIVATE": "PRIVATE_SUPERNET",
        "PRIVATE_RANGE": "PRIVATE_SUPERNET",
        "PUBLIC_RANGE": "PUBLIC",
    }
    explicit = aliases.get(explicit, explicit)
    if explicit in {"PUBLIC", "PRIVATE_SUPERNET", "VLAN"}:
        return explicit
    if row.get("vlan_name") not in (None, "") or row.get("vlan_tag") not in (
        None,
        "",
    ):
        return "VLAN"
    return "PRIVATE_SUPERNET" if network.is_private else "PUBLIC"


def _derive_site_code(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_").upper()


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value).lower().replace("_", " "))


def _tags(value: Any) -> list[str]:
    if isinstance(value, str):
        return [item.strip() for item in value.split(",") if item.strip()]
    if isinstance(value, (list, tuple)):
        return [str(item).strip() for item in value if str(item).strip()]
    return []


def _optional(value: Any) -> str | None:
    return _text(value) or None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()
