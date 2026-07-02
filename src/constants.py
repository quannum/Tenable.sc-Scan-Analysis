from datetime import datetime
from pathlib import Path

from openpyxl.styles import Font, PatternFill

VERSION = "2.2"


def default_output_file(run_datetime: datetime | None = None) -> Path:
    timestamp = (run_datetime or datetime.now()).strftime("%Y%m%d-%H%M%S")
    return Path("output") / f"tenable_scan_summary-{timestamp}.xlsx"


INCLUDE = "Include"
EXCLUDE = "Exclude"

STATUS_OK = "OK"
STATUS_PARTIAL = "PARTIAL"
STATUS_GAP = "GAP"

DEFAULT_EXPECTED_SHEET = "rsg-all"
FALLBACK_EXPECTED_SHEET = "Expected_Ranges"

SHEET_SCAN_SCOPE_SUMMARY = "Scan_Scope_Summary"
SHEET_SCAN_SCOPE_NORMALIZED = "Scan_Scope_Normalized"
SHEET_EXPECTED_VS_ACTUAL = "Expected_vs_Actual"
SHEET_EXPECTED_RANGE_COMPLIANCE = "Expected_Range_Compliance"
SHEET_TOP_EXCLUSION_IMPACT = "Top_Exclusion_Impact_Scans"
SHEET_EXECUTIVE_SUMMARY = "Executive_Summary"
SHEET_WARNINGS = "Warnings"
SHEET_RUN_METADATA = "Run_Metadata"

HEADER_FONT = Font(bold=True)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="F4CCCC")
GRAY = PatternFill("solid", fgColor="E7E6E6")
