import json
import shutil
import unittest
import importlib
from datetime import datetime
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

main = importlib.import_module("tenable_scan_analysis.cli.main")
from tenable_scan_analysis.io.app_config import Config
from tenable_scan_analysis.constants import (
    SHEET_EXECUTIVE_SUMMARY,
    SHEET_EXPECTED_VS_ACTUAL,
    SHEET_RUN_METADATA,
    SHEET_WARNINGS,
)


class EndToEndTests(unittest.TestCase):
    def test_pipeline_generates_expected_workbook_and_warning_sheet(self):
        temp_root = Path.cwd() / ".tmp-test-artifacts"
        temp_path = temp_root / "end_to_end_case"
        if temp_path.exists():
            shutil.rmtree(temp_path)

        temp_path.mkdir(parents=True, exist_ok=True)

        try:
            scan_dir = temp_path / "scans"
            asset_dir = temp_path / "assets"
            output_file = temp_path / "output" / "report.xlsx"
            expected_file = temp_path / "expected.xlsx"

            scan_dir.mkdir()
            asset_dir.mkdir()

            scan_payload = {
                "id": 1,
                "name": "Weekly Network Scan",
                "ipList": "10.0.0.0/24",
                "assets": [{"id": 100}, {"id": 999}],
                "schedule": {"enabled": True},
            }
            asset_payload = {
                "id": 100,
                "name": "Server Segment",
                "type": "static",
                "typeFields": {"definedIPs": "10.0.1.0/24"},
            }

            (scan_dir / "1_scan.json").write_text(json.dumps(scan_payload), encoding="utf-8")
            (asset_dir / "100_asset.json").write_text(json.dumps(asset_payload), encoding="utf-8")
            (asset_dir / "broken.json").write_text("{not valid json", encoding="utf-8")

            expected_wb = Workbook()
            expected_ws = expected_wb.active
            expected_ws.title = "rsg-all"
            expected_ws.append(["Scope Item", "Location", "Environment", "Required Scan"])
            expected_ws.append(["10.0.0.0/24", "HQ", "Prod", "Weekly Network Scan"])
            expected_ws.append(["10.0.1.0/25", "HQ", "Prod", "Missing Scan"])
            expected_wb.save(expected_file)

            config = Config(
                mode="offline",
                scan_json_dir=str(scan_dir),
                asset_json_dir=str(asset_dir),
                expected_scope_file=str(expected_file),
                expected_sheet=None,
                output_file=output_file,
                sc_access_key=None,
                sc_secret_key=None,
                sc_url=None,
                include_keywords=[],
                exclude_keywords=[],
                match_all_include=False,
                case_sensitive=False,
                filter_disabled_mode="ALL",
                log_level="INFO",
                log_file=None,
                run_started_at=datetime(2026, 5, 7, 13, 2, 3),
            )

            with patch("sys.stderr", new=StringIO()):
                collector = main.configure_logging("INFO")
                result_path = main.run_analysis(config, warning_records=collector.records)

            self.assertEqual(result_path, output_file)
            self.assertTrue(output_file.exists())

            workbook = load_workbook(output_file)
            self.assertIn(SHEET_EXPECTED_VS_ACTUAL, workbook.sheetnames)
            self.assertIn(SHEET_EXECUTIVE_SUMMARY, workbook.sheetnames)
            self.assertIn(SHEET_WARNINGS, workbook.sheetnames)
            self.assertIn(SHEET_RUN_METADATA, workbook.sheetnames)

            compare_ws = workbook[SHEET_EXPECTED_VS_ACTUAL]
            self.assertEqual(compare_ws["J2"].value, "OK")
            self.assertEqual(compare_ws["J3"].value, "OK")
            self.assertEqual(compare_ws["L2"].value, "Weekly Network Scan")
            self.assertEqual(compare_ws["M2"].value, "Yes")
            self.assertEqual(compare_ws["L3"].value, "Missing Scan")
            self.assertEqual(compare_ws["M3"].value, "No")

            exec_ws = workbook[SHEET_EXECUTIVE_SUMMARY]
            self.assertEqual(exec_ws["A2"].value, "Total Expected IPs")
            self.assertEqual(exec_ws["B2"].value, 384)
            self.assertEqual(exec_ws["B3"].value, 384)
            self.assertEqual(exec_ws["B4"].value, 0)

            warning_ws = workbook[SHEET_WARNINGS]
            warning_messages = [row[2] for row in warning_ws.iter_rows(min_row=2, values_only=True)]
            self.assertTrue(
                any("Skipping unreadable JSON file" in message for message in warning_messages)
            )
            self.assertTrue(
                any("references missing asset id '999'" in message for message in warning_messages)
            )

            metadata_ws = workbook[SHEET_RUN_METADATA]
            metadata = {
                row[0]: row[1]
                for row in metadata_ws.iter_rows(min_row=2, values_only=True)
            }
            self.assertEqual(metadata["Run Started At"], "2026-05-07 13:02:03")
            self.assertEqual(metadata["Mode"], "offline")
            self.assertEqual(metadata["Output File"], str(output_file))
            self.assertIsNone(metadata["Log File"])
        finally:
            if temp_path.exists():
                shutil.rmtree(temp_path)


if __name__ == "__main__":
    unittest.main()
