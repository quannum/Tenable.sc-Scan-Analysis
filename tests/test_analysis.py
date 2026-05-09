import unittest
from collections import defaultdict

from openpyxl import Workbook

from src.core.analysis import (
    ExcludedScopeRecord,
    ScopeRecord,
    calculate_coverage_result,
    calculate_scan_intervals,
    resolve_expected_sheet,
)
from src.core.scope_utils import parse_scope_item, scope_to_interval


class AnalysisTests(unittest.TestCase):
    def test_required_scan_matching_is_reported(self):
        actual = ScopeRecord(
            parsed=parse_scope_item("10.0.0.0/24"),
            scan_name="Required Weekly Scan",
            scope_item="10.0.0.0/24",
        )
        actual_by_scan = defaultdict(list, {actual.scan_name: [actual]})
        excluded_by_scan = defaultdict(list)
        exclusion_impact_by_scan = defaultdict(int)

        result = calculate_coverage_result(
            scope_item="10.0.0.0/24",
            location="HQ",
            environment="Prod",
            required_scan="Required Weekly Scan",
            expected=parse_scope_item("10.0.0.0/24"),
            actual_scopes=[actual],
            actual_by_scan=actual_by_scan,
            excluded_by_scan=excluded_by_scan,
            exclusion_impact_by_scan=exclusion_impact_by_scan,
        )

        self.assertEqual(result.required_scan, "Required Weekly Scan")
        self.assertEqual(result.required_scan_covered, "Yes")

    def test_required_scan_missing_is_reported(self):
        actual = ScopeRecord(
            parsed=parse_scope_item("10.0.0.0/24"),
            scan_name="Different Scan",
            scope_item="10.0.0.0/24",
        )
        actual_by_scan = defaultdict(list, {actual.scan_name: [actual]})
        excluded_by_scan = defaultdict(list)
        exclusion_impact_by_scan = defaultdict(int)

        result = calculate_coverage_result(
            scope_item="10.0.0.0/24",
            location="HQ",
            environment="Prod",
            required_scan="Required Weekly Scan",
            expected=parse_scope_item("10.0.0.0/24"),
            actual_scopes=[actual],
            actual_by_scan=actual_by_scan,
            excluded_by_scan=excluded_by_scan,
            exclusion_impact_by_scan=exclusion_impact_by_scan,
        )

        self.assertEqual(result.required_scan, "Required Weekly Scan")
        self.assertEqual(result.required_scan_covered, "No")

    def test_calculate_scan_intervals_subtracts_exclusions(self):
        scan_name = "Scan A"
        expected = parse_scope_item("10.0.0.1-10.0.0.10")
        expected_start, expected_end = scope_to_interval(expected)
        actual = ScopeRecord(
            parsed=parse_scope_item("10.0.0.1-10.0.0.10"),
            scan_name=scan_name,
            scope_item="10.0.0.1-10.0.0.10",
        )
        excluded = ExcludedScopeRecord(
            parsed=parse_scope_item("10.0.0.4-10.0.0.6"),
            scan_name=scan_name,
            asset_name="Excluded Segment",
            scope_item="10.0.0.4-10.0.0.6",
        )

        included_total, net_intervals, exclusions, excluded_total = (
            calculate_scan_intervals(
                scan_name=scan_name,
                expected=expected,
                expected_start=expected_start,
                expected_end=expected_end,
                actual_by_scan=defaultdict(list, {scan_name: [actual]}),
                excluded_by_scan=defaultdict(list, {scan_name: [excluded]}),
            )
        )

        net_total = sum(end - start + 1 for start, end in net_intervals)
        self.assertEqual(included_total, 10)
        self.assertEqual(excluded_total, 3)
        self.assertEqual(net_total, 7)
        self.assertEqual(exclusions[0].asset, "Excluded Segment")

    def test_resolve_expected_sheet_uses_requested_sheet(self):
        workbook = Workbook()
        workbook.active.title = "Default"
        workbook.create_sheet("CustomSheet")

        sheet = resolve_expected_sheet(workbook, expected_sheet="CustomSheet")

        self.assertEqual(sheet.title, "CustomSheet")

    def test_resolve_expected_sheet_matches_case_insensitively(self):
        workbook = Workbook()
        workbook.active.title = "Default"
        workbook.create_sheet("CustomSheet")

        sheet = resolve_expected_sheet(workbook, expected_sheet="customsheet")

        self.assertEqual(sheet.title, "CustomSheet")

    def test_resolve_expected_sheet_rejects_missing_requested_sheet(self):
        workbook = Workbook()
        workbook.active.title = "Default"

        with self.assertRaises(KeyError):
            resolve_expected_sheet(workbook, expected_sheet="MissingSheet")


if __name__ == "__main__":
    unittest.main()
